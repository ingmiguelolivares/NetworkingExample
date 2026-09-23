from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "GUIA_LAB_UNITY_NETWORKING_LINUX.md"
OUTPUT = ROOT / "GUIA_LAB_UNITY_NETWORKING_LINUX.docx"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
LIGHT_BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), LIGHT_BORDER)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_toc(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Abra el documento en Word y seleccione Actualizar tabla para mostrar la paginación."
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text):
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt(9.2)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def keep_lines(paragraph):
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = True


def add_code_block(document, lines):
    paragraph = document.add_paragraph(style="Code Block")
    keep_lines(paragraph)
    for idx, line in enumerate(lines):
        if idx:
            paragraph.add_run("\n")
        paragraph.add_run(line)
    return paragraph


def split_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(document, rows):
    data = [split_table_row(row) for row in rows]
    if len(data) > 1 and all(re.fullmatch(r":?-{3,}:?", value.replace(" ", "")) for value in data[1]):
        data.pop(1)
    cols = max(len(row) for row in data)
    table = document.add_table(rows=len(data), cols=cols)
    table.autofit = True
    table.style = "Table Grid"
    set_table_borders(table)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_idx, values in enumerate(data):
        row_pr = table.rows[row_idx]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)
        for col_idx in range(cols):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, NAVY)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if col_idx < len(values):
                add_inline(paragraph, values[col_idx])
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if row_idx == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def setup_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 18, 18, 7), ("Heading 2", 14, 14, 5), ("Heading 3", 11.5, 10, 3)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Aptos Mono"
    code.font.size = Pt(8.6)
    code.font.color.rgb = RGBColor(25, 25, 25)
    code.paragraph_format.left_indent = Cm(0.55)
    code.paragraph_format.right_indent = Cm(0.35)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.0
    p_pr = code.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GRAY)
    p_pr.append(shd)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.2)
        style.paragraph_format.space_after = Pt(3)


def add_cover(document):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(70)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("GUÍA DE LABORATORIO")
    run.font.name = "Aptos Display"
    run.font.size = Pt(16)
    run.bold = True

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Unity Multiplayer Networking")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("Cliente Servidor con Linux Kubernetes y Agones")
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.bold = True
    stack = document.add_paragraph()
    stack.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stack.paragraph_format.space_after = Pt(70)
    run = stack.add_run("Netcode for GameObjects y Netcode for Entities")
    run.font.name = "Aptos"
    run.font.size = Pt(13)

    fields = document.add_table(rows=4, cols=2)
    fields.autofit = False
    fields.columns[0].width = Cm(4)
    fields.columns[1].width = Cm(11.5)
    set_table_borders(fields)
    for idx, label in enumerate(("Asignatura", "Docente", "Estudiante", "Fecha")):
        fields.cell(idx, 0).text = label
        fields.cell(idx, 1).text = ""
        set_cell_margins(fields.cell(idx, 0), 120, 140, 120, 140)
        set_cell_margins(fields.cell(idx, 1), 120, 140, 120, 140)
        fields.cell(idx, 0).paragraphs[0].runs[0].bold = True

    document.add_paragraph().paragraph_format.space_after = Pt(28)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("NetworkingExample Starter Project\nUnity 6000.6.0f1")
    run.font.size = Pt(10)
    document.add_page_break()


def parse_markdown(document, markdown):
    lines = markdown.splitlines()
    index = 0
    skipped_titles = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        if stripped.startswith("# ") or stripped.startswith("## "):
            if skipped_titles < 2:
                skipped_titles += 1
                index += 1
                continue
            heading = stripped.lstrip("#").strip().replace("—", " ")
            paragraph = document.add_paragraph(heading, style="Heading 1")
            keep_with_next(paragraph)
            index += 1
            continue
        if stripped.startswith("### "):
            heading = stripped[4:].replace("`", "").replace("—", " ")
            paragraph = document.add_paragraph(heading, style="Heading 2")
            keep_with_next(paragraph)
            index += 1
            continue
        if stripped.startswith("#### "):
            heading = stripped[5:].replace("`", "").replace("—", " ")
            paragraph = document.add_paragraph(heading, style="Heading 3")
            keep_with_next(paragraph)
            index += 1
            continue
        if re.fullmatch(r"=+", stripped):
            index += 1
            continue
        if stripped.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                value = lines[index].strip().lstrip(">").strip()
                value = re.sub(r"^\[!(IMPORTANT|NOTE|WARNING|CAUTION)\]\s*", lambda m: m.group(1).title() + ". ", value)
                quote_lines.append(value)
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.55)
            paragraph.paragraph_format.right_indent = Cm(0.35)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline(paragraph, " ".join(quote_lines))
            if paragraph.runs:
                paragraph.runs[0].bold = True
            keep_lines(paragraph)
            continue
        list_match = re.match(r"^[-*]\s+(.*)", stripped)
        numbered_match = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if list_match or numbered_match:
            if list_match:
                paragraph = document.add_paragraph(style="List Bullet")
                content = list_match.group(1)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.65)
                paragraph.paragraph_format.first_line_indent = Cm(-0.45)
                paragraph.paragraph_format.space_after = Pt(3)
                content = f"{numbered_match.group(1)}. {numbered_match.group(2)}"
            add_inline(paragraph, content)
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "```", "|", ">", "- ", "* ")) or re.match(r"^\d+\.\s+", candidate) or re.fullmatch(r"=+", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines).replace("  ", " "))
        keep_lines(paragraph)


def main():
    markdown = SOURCE.read_text(encoding="utf-8")
    document = Document()
    setup_styles(document)
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header.paragraphs[0]
    header.text = "Unity Multiplayer Networking  Guía de laboratorio"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(80, 80, 80)
    add_page_field(section.footer.paragraphs[0])

    add_cover(document)
    toc_title = document.add_paragraph("Tabla de contenido", style="Heading 1")
    toc_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_toc(document.add_paragraph())
    outline = document.add_paragraph()
    outline.add_run("Secciones principales").bold = True
    for line in markdown.splitlines():
        if re.match(r"^## \d+\.", line):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.add_run(line[3:].replace("—", " "))
    document.add_page_break()
    parse_markdown(document, markdown)

    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    document.core_properties.title = "Guía de laboratorio Unity Multiplayer Networking"
    document.core_properties.subject = "Starter Project NGO NFE Linux Kubernetes Agones"
    document.core_properties.author = ""
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
